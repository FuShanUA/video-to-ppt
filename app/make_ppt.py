#!/usr/bin/env python3
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu


IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".heic",
    ".tif",
    ".tiff",
    ".webp",
    ".bmp",
}

VIDEO_EXTENSIONS = {
    ".mp4",
    ".mov",
    ".m4v",
    ".mkv",
    ".avi",
    ".webm",
    ".flv",
    ".wmv",
}


def natural_key(path: Path) -> tuple[tuple[int, object], ...]:
    parts = re.split(r"(\d+)", path.name)
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in parts
        if part
    )


def list_images(folder: Path) -> list[Path]:
    if not folder.is_dir():
        raise ValueError(f"不是文件夹：{folder}")

    return sorted(
        (
            path
            for path in folder.iterdir()
            if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS
        ),
        key=natural_key,
    )


def build_presentation(
    folder: Path, output_path: Path | None = None, fit_mode: str = "fit"
) -> Path:
    images = list_images(folder)
    if not images:
        raise ValueError("文件夹里没有找到图片。")
    if fit_mode not in {"fit", "cover"}:
        raise ValueError("图片排版模式必须是 fit 或 cover。")

    presentation = Presentation()
    presentation.slide_width = Emu(12192000)
    presentation.slide_height = Emu(6858000)
    blank_layout = presentation.slide_layouts[6]
    target_aspect = presentation.slide_width / presentation.slide_height

    for image_path in images:
        with Image.open(image_path) as image:
            source_aspect = image.width / image.height

        slide = presentation.slides.add_slide(blank_layout)

        if fit_mode == "fit":
            if source_aspect > target_aspect:
                width = presentation.slide_width
                height = int(width / source_aspect)
                left = 0
                top = int((presentation.slide_height - height) / 2)
            else:
                height = presentation.slide_height
                width = int(height * source_aspect)
                left = int((presentation.slide_width - width) / 2)
                top = 0

            slide.shapes.add_picture(
                str(image_path),
                left,
                top,
                width=width,
                height=height,
            )
            continue

        picture = slide.shapes.add_picture(
            str(image_path),
            0,
            0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )

        if source_aspect < target_aspect:
            crop = (1 - source_aspect / target_aspect) / 2
            picture.crop_top = crop
            picture.crop_bottom = crop
        elif source_aspect > target_aspect:
            crop = (1 - target_aspect / source_aspect) / 2
            picture.crop_left = crop
            picture.crop_right = crop

    output_path = output_path or folder.parent / f"{folder.name}_16x9.pptx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    return output_path


def build_video_presentation(video: Path) -> Path:
    helper = Path(__file__).with_name("video_to_ppt.py")
    output_path = video.with_suffix(".pptx")
    result = subprocess.run(
        [
            sys.executable,
            str(helper),
            str(video),
            "--output",
            str(output_path),
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    if result.returncode != 0:
        message = result.stderr.strip() or result.stdout.strip() or "未知错误"
        raise ValueError(message)
    if not output_path.is_file():
        raise ValueError("视频处理结束，但没有找到输出文件。")
    return output_path


def main() -> int:
    if len(sys.argv) != 2:
        print("用法：make_ppt.py <图片文件夹或视频文件>", file=sys.stderr)
        return 2

    try:
        target = Path(sys.argv[1]).expanduser()
        if target.is_dir():
            output_path = build_presentation(target)
        elif target.is_file() and target.suffix.lower() in VIDEO_EXTENSIONS:
            output_path = build_video_presentation(target)
        else:
            raise ValueError("请拖入图片文件夹，或 mp4/mov/mkv/webm 等视频文件。")
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        return 1

    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

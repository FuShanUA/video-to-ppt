from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation

from app.make_ppt import build_presentation, list_images, natural_key


class MakePptTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def create_image(self, name: str, width: int, height: int) -> Path:
        path = self.folder / name
        Image.new("RGB", (width, height), "#336699").save(path)
        return path

    def test_natural_order(self) -> None:
        paths = [
            self.create_image("slide10.png", 10, 10),
            self.create_image("slide2.png", 10, 10),
            self.create_image("slide1.png", 10, 10),
        ]
        self.assertEqual([path.name for path in sorted(paths, key=natural_key)], [
            "slide1.png",
            "slide2.png",
            "slide10.png",
        ])

    def test_fit_mode_keeps_the_whole_image(self) -> None:
        self.create_image("portrait.png", 800, 1600)
        output = self.folder / "output.pptx"

        build_presentation(self.folder, output, fit_mode="fit")

        presentation = Presentation(output)
        slide = presentation.slides[0]
        picture = slide.shapes[0]
        slide_width = presentation.slide_width
        slide_height = presentation.slide_height
        self.assertEqual(picture.width, slide_height // 2)
        self.assertEqual(picture.height, slide_height)
        self.assertEqual(picture.left, (slide_width - picture.width) // 2)
        self.assertEqual(picture.top, 0)
        self.assertEqual(picture.crop_top, 0)
        self.assertEqual(picture.crop_bottom, 0)

    def test_cover_mode_crops_to_slide(self) -> None:
        self.create_image("portrait.png", 800, 1600)
        output = self.folder / "output.pptx"

        build_presentation(self.folder, output, fit_mode="cover")

        presentation = Presentation(output)
        picture = presentation.slides[0].shapes[0]
        self.assertEqual(picture.width, presentation.slide_width)
        self.assertEqual(picture.height, presentation.slide_height)
        self.assertGreater(picture.crop_top, 0)
        self.assertGreater(picture.crop_bottom, 0)

    def test_invalid_fit_mode(self) -> None:
        self.create_image("image.png", 100, 100)
        with self.assertRaisesRegex(ValueError, "fit 或 cover"):
            build_presentation(self.folder, fit_mode="stretch")

    def test_list_images_only_accepts_directories(self) -> None:
        with self.assertRaisesRegex(ValueError, "不是文件夹"):
            list_images(self.folder / "missing")


if __name__ == "__main__":
    unittest.main()

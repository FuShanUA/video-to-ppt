#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import subprocess
import sys
import traceback
from dataclasses import dataclass, asdict
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu


def find_executable(name: str) -> str:
    found = shutil.which(name)
    if found:
        return found
    for directory in ("/opt/homebrew/bin", "/usr/local/bin", "/usr/bin", "/bin"):
        candidate = Path(directory) / name
        if candidate.is_file() and candidate.is_executable():
            return str(candidate)
    return name


FFMPEG = find_executable("ffmpeg")
FFPROBE = find_executable("ffprobe")


@dataclass
class Candidate:
    index: int
    start: int
    end: int
    representative: int
    duration: float
    motion: float
    forced: bool = False
    sharpness: float = 0.0
    border_uniformity: float = 0.0
    border_edge_density: float = 0.0
    content_edge_density: float = 0.0
    software_score: int = 0
    action: str = "kept"
    reason: str = ""


def run_command(command: list[str]) -> None:
    result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if result.returncode != 0:
        message = result.stderr.strip() or result.stdout.strip() or "unknown FFmpeg error"
        raise RuntimeError(message)


def probe_video(video: Path) -> dict:
    command = [
        FFPROBE, "-v", "error", "-select_streams", "v:0",
        "-show_entries", "stream=width,height,avg_frame_rate",
        "-show_entries", "format=duration",
        "-of", "json", str(video),
    ]
    result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, check=True)
    payload = json.loads(result.stdout)
    streams = payload.get("streams") or []
    if not streams:
        raise ValueError(f"没有找到视频流：{video}")
    stream = streams[0]
    rate_text = stream.get("avg_frame_rate") or "0/0"
    try:
        numerator, denominator = rate_text.split("/", 1)
        frame_rate = float(numerator) / float(denominator) if float(denominator) else 0.0
    except ValueError:
        frame_rate = 0.0
    if not frame_rate or not math.isfinite(frame_rate):
        frame_rate = 25.0
    return {
        "width": int(stream["width"]),
        "height": int(stream["height"]),
        "frame_rate": frame_rate,
        "duration": float(payload.get("format", {}).get("duration", 0.0)),
    }


def parse_crop(text: str) -> tuple[int, int, int, int]:
    parts = [part.strip() for part in text.split(",")]
    if len(parts) != 4:
        raise argparse.ArgumentTypeError("crop 格式应为 x,y,w,h")
    try:
        x, y, w, h = map(int, parts)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("crop 四个值必须是整数") from exc
    if w <= 0 or h <= 0:
        raise argparse.ArgumentTypeError("crop 宽高必须大于 0")
    return x, y, w, h


def parse_time(text: str) -> float:
    value = text.strip()
    if re.fullmatch(r"\d+(?:\.\d+)?", value):
        return float(value)
    parts = value.split(":")
    if len(parts) > 3 or not parts:
        raise argparse.ArgumentTypeError("时间格式应为秒、MM:SS 或 HH:MM:SS")
    try:
        numbers = [float(part) for part in parts]
    except ValueError as exc:
        raise argparse.ArgumentTypeError("时间格式应为秒、MM:SS 或 HH:MM:SS") from exc
    while len(numbers) < 3:
        numbers.insert(0, 0.0)
    hours, minutes, seconds = numbers
    if minutes >= 60 or seconds >= 60:
        raise argparse.ArgumentTypeError("分钟和秒必须小于 60")
    return hours * 3600 + minutes * 60 + seconds


def format_time(seconds: float) -> str:
    seconds = max(0, int(round(seconds)))
    hours, remainder = divmod(seconds, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def largest_run(values: np.ndarray) -> tuple[int, int]:
    if len(values) == 0:
        return 0, -1
    indexes = np.flatnonzero(values)
    if len(indexes) == 0:
        return 0, -1
    best = (int(indexes[0]), int(indexes[0]))
    current = best
    for previous, current_index in zip(indexes, indexes[1:]):
        if int(current_index) == int(previous) + 1:
            current = (current[0], int(current_index))
        else:
            best = max(best, current, key=lambda item: item[1] - item[0])
            current = (int(current_index), int(current_index))
    return max(best, current, key=lambda item: item[1] - item[0])


def clear_generated(workdir: Path) -> None:
    for name in ("frames", "slides", "review"):
        target = workdir / name
        if target.exists():
            shutil.rmtree(target)
    for name in ("manifest.tsv", "report.json"):
        target = workdir / name
        if target.exists():
            target.unlink()


def extract_lowres(video: Path, workdir: Path, fps: float, resume: bool) -> list[Path]:
    output_dir = workdir / "frames"
    output_dir.mkdir(parents=True, exist_ok=True)
    frames = sorted(output_dir.glob("frame_*.jpg"))
    if resume and frames:
        return frames
    for frame in frames:
        frame.unlink()
    command = [
        FFMPEG, "-hide_banner", "-loglevel", "error", "-nostdin",
        "-i", str(video),
        "-vf", f"fps={fps:g},scale=320:-2:flags=bilinear",
        "-vsync", "0", "-q:v", "4",
        str(output_dir / "frame_%06d.jpg"),
    ]
    run_command(command)
    frames = sorted(output_dir.glob("frame_*.jpg"))
    if not frames:
        raise RuntimeError("抽帧结果为空")
    return frames


def detect_crop(
    frames: list[Path], source_width: int, source_height: int
) -> tuple[tuple[int, int, int, int], tuple[int, int, int, int], str]:
    lowres_width, lowres_height = Image.open(frames[0]).size
    sample_count = min(240, len(frames))
    step = max(1, len(frames) // sample_count)
    sampled = frames[::step]
    stack = np.stack([
        np.asarray(Image.open(path).convert("L"), dtype=np.float32)
        for path in sampled
    ])
    temporal_std = stack.std(axis=0)
    row_activity = temporal_std.mean(axis=1)
    column_activity = temporal_std.mean(axis=0)

    if row_activity.max() < 2.0 or column_activity.max() < 2.0:
        source_crop = (0, 0, source_width, source_height)
        lowres_crop = (0, 0, lowres_width, lowres_height)
        return source_crop, lowres_crop, "full-frame: low temporal variance"

    row_start, row_end = largest_run(row_activity >= row_activity.max() * 0.2)
    col_start, col_end = largest_run(column_activity >= column_activity.max() * 0.2)
    row_coverage = (row_end - row_start + 1) / lowres_height
    col_coverage = (col_end - col_start + 1) / lowres_width
    if row_coverage > 0.90 and col_coverage > 0.90:
        source_crop = (0, 0, source_width, source_height)
        lowres_crop = (0, 0, lowres_width, lowres_height)
        return source_crop, lowres_crop, "full-frame: high-variance area covers frame"

    row_start = max(0, row_start - max(1, round(lowres_height * 0.10)))
    scale_x = source_width / lowres_width
    scale_y = source_height / lowres_height
    x = max(0, round(col_start * scale_x))
    y = max(0, round(row_start * scale_y))
    width = min(source_width - x, max(1, round((col_end - col_start + 1) * scale_x)))
    height = min(source_height - y, max(1, round((row_end - row_start + 1) * scale_y)))
    source_crop = (x, y, width, height)

    low_x = max(0, col_start)
    low_y = max(0, row_start)
    low_w = min(lowres_width - low_x, col_end - col_start + 1)
    low_h = min(lowres_height - low_y, row_end - row_start + 1)
    lowres_crop = (low_x, low_y, low_w, low_h)
    return source_crop, lowres_crop, "temporal-variance bounding region"


def crop_from_source(
    crop: tuple[int, int, int, int], source_width: int, source_height: int,
    lowres_width: int, lowres_height: int
) -> tuple[int, int, int, int]:
    x, y, w, h = crop
    scale_x = lowres_width / source_width
    scale_y = lowres_height / source_height
    low_x = max(0, round(x * scale_x))
    low_y = max(0, round(y * scale_y))
    low_w = max(1, min(lowres_width - low_x, round(w * scale_x)))
    low_h = max(1, min(lowres_height - low_y, round(h * scale_y)))
    return low_x, low_y, low_w, low_h


def dhash(image: Image.Image) -> int:
    small = image.resize((9, 8), Image.Resampling.BILINEAR)
    pixels = np.asarray(small, dtype=np.int16)
    bits = 0
    for row in range(8):
        for col in range(8):
            left = int(pixels[row, col])
            right = int(pixels[row, col + 1])
            bits = (bits << 1) | (1 if left > right else 0)
    return bits


def load_feature(path: Path, crop: tuple[int, int, int, int]) -> tuple[np.ndarray, int, np.ndarray]:
    with Image.open(path) as image:
        region = image.convert("L").crop(crop)
        feature = np.asarray(region.resize((64, 40), Image.Resampling.BILINEAR), dtype=np.int16)
        hash_value = dhash(region)
        rgb = np.asarray(region.resize((32, 20), Image.Resampling.BILINEAR), dtype=np.uint8)
    quantized = (rgb // 64).astype(np.int16)
    flat = quantized[..., 0] * 16 + quantized[..., 1] * 4 + quantized[..., 2]
    histogram = np.bincount(flat.reshape(-1), minlength=64).astype(np.float64)
    histogram /= max(1.0, histogram.sum())
    return feature, hash_value, histogram


def mean_abs_diff(a: np.ndarray, b: np.ndarray) -> float:
    return float(np.abs(a - b).mean())


def histogram_correlation(a: np.ndarray, b: np.ndarray) -> float:
    denominator = math.sqrt(float(np.dot(a, a) * np.dot(b, b)))
    return float(np.dot(a, b)) / denominator if denominator else 0.0


def hamming_distance(a: int, b: int) -> int:
    return (a ^ b).bit_count()


def same_page(a: tuple[np.ndarray, int, np.ndarray], b: tuple[np.ndarray, int, np.ndarray]) -> bool:
    image_diff = mean_abs_diff(a[0], b[0])
    hash_diff = hamming_distance(a[1], b[1])
    histogram = histogram_correlation(a[2], b[2])
    return (
        (image_diff < 0.8 and hash_diff <= 8)
        or (image_diff < 1.5 and hash_diff <= 5)
        or (histogram > 0.998 and hash_diff <= 10)
    )


def build_runs(
    frames: list[Path], crop: tuple[int, int, int, int], fps: float,
    change_threshold: float, min_run: float
) -> tuple[list[Candidate], list[tuple[np.ndarray, int, np.ndarray]], list[float]]:
    features = [load_feature(path, crop) for path in frames]
    diffs = [0.0]
    for previous, current in zip(features, features[1:]):
        diffs.append(mean_abs_diff(previous[0], current[0]))

    runs: list[tuple[int, int]] = []
    start = 0
    for index in range(1, len(frames)):
        if diffs[index] >= change_threshold:
            runs.append((start, index - 1))
            start = index
    runs.append((start, len(frames) - 1))

    candidates: list[Candidate] = []
    for start_index, end_index in runs:
        duration = (end_index - start_index + 1) / fps
        if duration + 1e-9 < min_run:
            continue
        representative = (start_index + end_index) // 2
        motion = float(np.mean(diffs[start_index + 1:end_index + 1])) if end_index > start_index else 0.0
        candidates.append(Candidate(
            index=representative,
            start=start_index,
            end=end_index,
            representative=representative,
            duration=duration,
            motion=motion,
        ))
    return candidates, features, diffs


def add_forced_times(
    candidates: list[Candidate], frames: list[Path], fps: float, forced_times: list[float]
) -> list[Candidate]:
    if not forced_times:
        return candidates
    existing = {candidate.index for candidate in candidates}
    for seconds in forced_times:
        index = max(0, min(len(frames) - 1, round(seconds * fps)))
        if index in existing:
            for candidate in candidates:
                if candidate.index == index:
                    candidate.forced = True
            continue
        candidates.append(Candidate(
            index=index, start=index, end=index, representative=index,
            duration=1.0 / fps, motion=0.0, forced=True,
        ))
        existing.add(index)
    candidates.sort(key=lambda item: item.index)
    return candidates


def validate_chronological_order(candidates: list[Candidate]) -> None:
    indexes = [candidate.index for candidate in candidates]
    if any(current <= previous for previous, current in zip(indexes, indexes[1:])):
        raise RuntimeError("页面时间顺序校验失败：输出页序不是严格递增")


def layout_metrics(path: Path, crop: tuple[int, int, int, int]) -> dict[str, float]:
    with Image.open(path) as image:
        region = image.convert("RGB").crop(crop)
        gray = np.asarray(region.convert("L"), dtype=np.float32)
        rgb = np.asarray(region, dtype=np.uint8)
    height, width = gray.shape
    gx = np.zeros_like(gray)
    gy = np.zeros_like(gray)
    gx[:, 1:-1] = gray[:, 2:] - gray[:, :-2]
    gy[1:-1, :] = gray[2:, :] - gray[:-2, :]
    edges = np.sqrt(gx * gx + gy * gy) > 18.0

    border_height = max(2, round(height * 0.05))
    border_width = max(2, round(width * 0.02))
    border = np.zeros((height, width), dtype=bool)
    border[:border_height, :] = True
    border[-border_height:, :] = True
    border[:, :border_width] = True
    border[:, -border_width:] = True
    center = ~border

    border_rgb = rgb[border]
    quantized = (border_rgb // 16).astype(np.int16)
    keys = quantized[:, 0] * 256 + quantized[:, 1] * 16 + quantized[:, 2]
    counts = np.bincount(keys, minlength=4096)
    border_uniformity = float(counts.max() / max(1, counts.sum()))

    laplacian = np.zeros_like(gray)
    laplacian[1:-1, 1:-1] = (
        4.0 * gray[1:-1, 1:-1]
        - gray[:-2, 1:-1] - gray[2:, 1:-1]
        - gray[1:-1, :-2] - gray[1:-1, 2:]
    )
    sharpness = float(laplacian[1:-1, 1:-1].var())
    return {
        "sharpness": sharpness,
        "border_uniformity": border_uniformity,
        "border_edge_density": float(edges[border].mean()),
        "content_edge_density": float(edges[center].mean()),
    }


def score_software(candidate: Candidate, metrics: dict[str, float], exclude_software: bool) -> None:
    candidate.sharpness = metrics["sharpness"]
    candidate.border_uniformity = metrics["border_uniformity"]
    candidate.border_edge_density = metrics["border_edge_density"]
    candidate.content_edge_density = metrics["content_edge_density"]
    score = 0
    if candidate.border_uniformity < 0.60:
        score += 1
    if candidate.border_uniformity < 0.40:
        score += 2
    if candidate.border_edge_density > 0.05:
        score += 1
    if candidate.border_edge_density > 0.10:
        score += 2
    if candidate.content_edge_density > 0.20:
        score += 1
    if candidate.motion > 0.75:
        score += 2
    elif candidate.motion > 0.35:
        score += 1
    if candidate.duration > 120 and candidate.motion > 0.25:
        score += 1
    candidate.software_score = score
    if not exclude_software:
        candidate.action = "kept"
        candidate.reason = "software exclusion disabled"


def select_candidates(
    candidates: list[Candidate], frames: list[Path], crop: tuple[int, int, int, int],
    features: list[tuple[np.ndarray, int, np.ndarray]], software_threshold: int,
    exclude_software: bool, dedupe: bool, transition_sharpness: float
) -> list[Candidate]:
    accepted: list[Candidate] = []
    for candidate in candidates:
        metrics = layout_metrics(frames[candidate.index], crop)
        score_software(candidate, metrics, exclude_software)
        feature = features[candidate.index]

        if candidate.forced:
            candidate.action = "kept"
            candidate.reason = "forced timestamp"
            accepted.append(candidate)
            continue

        if transition_sharpness > 0 and candidate.end == candidate.start and candidate.sharpness < transition_sharpness:
            candidate.action = "excluded_transition"
            candidate.reason = f"single-frame run below sharpness {transition_sharpness:g}"
            continue

        if exclude_software and candidate.software_score >= software_threshold:
            candidate.action = "excluded_software"
            candidate.reason = f"software score {candidate.software_score} >= {software_threshold}"
            continue

        if dedupe:
            duplicate = next(
                (previous for previous in accepted if same_page(feature, features[previous.index])),
                None,
            )
            if duplicate is not None:
                candidate.action = "merged_duplicate"
                candidate.reason = f"duplicate of frame {duplicate.index}"
                continue

        candidate.action = "kept"
        candidate.reason = "stable representative"
        accepted.append(candidate)
    return accepted


def extract_full_frames(
    video: Path, workdir: Path, accepted: list[Candidate], fps: float,
    source_frame_rate: float, crop: tuple[int, int, int, int]
) -> list[Path]:
    slides_dir = workdir / "slides"
    slides_dir.mkdir(parents=True, exist_ok=True)
    for path in slides_dir.glob("slide_*.png"):
        path.unlink()

    frame_numbers = sorted({
        max(0, round((candidate.index / fps) * source_frame_rate))
        for candidate in accepted
    })
    select_expr = "+".join(f"eq(n\\,{number})" for number in frame_numbers)
    crop_filter = f"crop={crop[2]}:{crop[3]}:{crop[0]}:{crop[1]}"
    video_filter = f"select='{select_expr}',{crop_filter}"
    command = [
        FFMPEG, "-hide_banner", "-loglevel", "error", "-nostdin",
        "-i", str(video), "-vf", video_filter, "-vsync", "0",
        "-frames:v", str(len(frame_numbers)), "-start_number", "1",
        str(slides_dir / "slide_%04d.png"),
    ]
    try:
        run_command(command)
    except RuntimeError:
        pass

    outputs = sorted(slides_dir.glob("slide_*.png"))
    if len(outputs) == len(frame_numbers):
        return outputs

    for path in slides_dir.glob("slide_*.png"):
        path.unlink()
    for page_number, candidate in enumerate(accepted, 1):
        seconds = max(0.0, candidate.index / fps - 0.5 / fps)
        command = [
            "ffmpeg", "-hide_banner", "-loglevel", "error", "-nostdin",
            "-ss", f"{seconds:.6f}", "-i", str(video),
            "-frames:v", "1", "-vf", crop_filter,
            str(slides_dir / f"slide_{page_number:04d}.png"),
        ]
        run_command(command)
    outputs = sorted(slides_dir.glob("slide_*.png"))
    if len(outputs) != len(accepted):
        raise RuntimeError(f"全帧提取数量不匹配：期望 {len(accepted)}，得到 {len(outputs)}")
    return outputs


def build_pptx(images: list[Path], output: Path, candidates: list[Candidate], fps: float) -> None:
    presentation = Presentation()
    presentation.slide_width = Emu(12192000)
    presentation.slide_height = Emu(6858000)
    blank_layout = presentation.slide_layouts[6]
    target_aspect = presentation.slide_width / presentation.slide_height

    for image_path, candidate in zip(images, candidates):
        with Image.open(image_path) as image:
            source_aspect = image.width / image.height
        slide = presentation.slides.add_slide(blank_layout)
        picture = slide.shapes.add_picture(
            str(image_path), 0, 0,
            width=presentation.slide_width, height=presentation.slide_height,
        )
        if source_aspect < target_aspect:
            crop = (1 - source_aspect / target_aspect) / 2
            picture.crop_top = crop
            picture.crop_bottom = crop
        elif source_aspect > target_aspect:
            crop = (1 - target_aspect / source_aspect) / 2
            picture.crop_left = crop
            picture.crop_right = crop
        notes = slide.notes_slide.notes_text_frame
        notes.text = (
            f"Source time: {format_time(candidate.index / fps)}\n"
            f"Stable run: {format_time(candidate.start / fps)} - {format_time(candidate.end / fps)}"
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)


def write_manifest(workdir: Path, candidates: list[Candidate], fps: float) -> None:
    columns = [
        "action", "reason", "source_index", "source_time", "run_start", "run_end",
        "duration_seconds", "motion", "forced", "software_score", "sharpness",
        "border_uniformity", "border_edge_density", "content_edge_density",
    ]
    with (workdir / "manifest.tsv").open("w", encoding="utf-8") as handle:
        handle.write("\t".join(columns) + "\n")
        for candidate in candidates:
            values = [
                candidate.action, candidate.reason, candidate.index,
                format_time(candidate.index / fps), format_time(candidate.start / fps),
                format_time(candidate.end / fps), f"{candidate.duration:.3f}",
                f"{candidate.motion:.4f}", str(int(candidate.forced)),
                str(candidate.software_score), f"{candidate.sharpness:.2f}",
                f"{candidate.border_uniformity:.4f}", f"{candidate.border_edge_density:.4f}",
                f"{candidate.content_edge_density:.4f}",
            ]
            handle.write("\t".join(map(str, values)) + "\n")


def make_contact_sheet(images: list[Path], candidates: list[Candidate], fps: float, output: Path) -> None:
    if not images:
        return
    output.parent.mkdir(parents=True, exist_ok=True)
    thumb_width = 320
    columns = 5
    rows = math.ceil(len(images) / columns)
    label_height = 28
    with Image.open(images[0]) as first:
        thumb_height = round(thumb_width * first.height / first.width)
    sheet = Image.new("RGB", (columns * thumb_width, rows * (thumb_height + label_height)), "white")
    draw = ImageDraw.Draw(sheet)
    font = ImageFont.load_default()
    for position, (image_path, candidate) in enumerate(zip(images, candidates)):
        row, column = divmod(position, columns)
        x = column * thumb_width
        y = row * (thumb_height + label_height)
        with Image.open(image_path) as image:
            thumbnail = image.convert("RGB").resize(
                (thumb_width, thumb_height), Image.Resampling.BILINEAR
            )
        sheet.paste(thumbnail, (x, y))
        label = f"{position + 1:03d}  {format_time(candidate.index / fps)}"
        draw.text((x + 6, y + thumb_height + 5), label, fill=(30, 30, 30), font=font)
    sheet.save(output, quality=90)


def main() -> int:
    parser = argparse.ArgumentParser(description="从视频中提取 PPT 页面并生成 16:9 PPTX")
    parser.add_argument("video", type=Path, help="输入视频")
    parser.add_argument("--output", type=Path, help="输出 PPTX，默认放在视频旁")
    parser.add_argument("--workdir", type=Path, help="工作目录，默认放在视频旁")
    parser.add_argument("--fps", type=float, default=2.0, help="抽帧频率，默认 2")
    parser.add_argument("--crop", type=parse_crop, help="源视频像素裁剪区域 x,y,w,h")
    parser.add_argument("--change-threshold", type=float, default=2.0, help="帧差阈值，默认 2.0")
    parser.add_argument("--min-run", type=float, default=0.5, help="最短稳定区间秒数，默认 0.5")
    parser.add_argument("--software-threshold", type=int, default=5, help="纯软件演示排除分数，默认 5")
    parser.add_argument("--transition-sharpness", type=float, default=0.0, help="单帧运行的最低清晰度，默认不启用")
    parser.add_argument("--force-time", action="append", default=[], type=parse_time, help="强制保留时间点，可重复")
    parser.add_argument("--no-software-exclusion", action="store_true", help="保留所有软件演示画面")
    parser.add_argument("--no-dedupe", action="store_true", help="关闭重复页合并")
    parser.add_argument("--resume", action="store_true", help="复用已有低分辨率帧")
    args = parser.parse_args()

    video = args.video.expanduser()
    if not video.is_file():
        print(f"视频不存在：{video}", file=sys.stderr)
        return 1
    if args.fps <= 0:
        print("--fps 必须大于 0", file=sys.stderr)
        return 1

    output = (args.output or video.with_suffix(".pptx")).expanduser()
    workdir = (args.workdir or video.parent / f".{video.stem}_video_ppt_work").expanduser()
    workdir.mkdir(parents=True, exist_ok=True)
    if not args.resume:
        clear_generated(workdir)

    try:
        video_info = probe_video(video)
        if args.crop:
            x, y, w, h = args.crop
            if x + w > video_info["width"] or y + h > video_info["height"]:
                raise ValueError("crop 超出视频范围")
            source_crop = args.crop
        frames = extract_lowres(video, workdir, args.fps, args.resume)
        lowres_width, lowres_height = Image.open(frames[0]).size
        if args.crop:
            lowres_crop = crop_from_source(
                source_crop, video_info["width"], video_info["height"],
                lowres_width, lowres_height,
            )
            crop_method = "manual"
        else:
            source_crop, lowres_crop, crop_method = detect_crop(
                frames, video_info["width"], video_info["height"]
            )

        candidates, features, _ = build_runs(
            frames, lowres_crop, args.fps, args.change_threshold, args.min_run
        )
        candidates = add_forced_times(candidates, frames, args.fps, args.force_time)
        accepted = select_candidates(
            candidates, frames, lowres_crop, features,
            args.software_threshold, not args.no_software_exclusion,
            not args.no_dedupe, args.transition_sharpness,
        )
        if not accepted:
            raise RuntimeError("没有可保留的页面；可降低 --change-threshold 或使用 --force-time")
        validate_chronological_order(accepted)

        slide_images = extract_full_frames(
            video, workdir, accepted, args.fps,
            video_info["frame_rate"], source_crop,
        )
        build_pptx(slide_images, output, accepted, args.fps)
        write_manifest(workdir, candidates, args.fps)
        make_contact_sheet(
            slide_images, accepted, args.fps, workdir / "review" / "contact_sheet.jpg"
        )

        report = {
            "video": str(video),
            "output": str(output),
            "workdir": str(workdir),
            "fps": args.fps,
            "source": video_info,
            "crop": source_crop,
            "crop_method": crop_method,
            "sampled_frames": len(frames),
            "stable_candidates": len(candidates),
            "kept_pages": len(accepted),
            "excluded_software": sum(c.action == "excluded_software" for c in candidates),
            "merged_duplicates": sum(c.action == "merged_duplicate" for c in candidates),
            "excluded_transitions": sum(c.action == "excluded_transition" for c in candidates),
            "forced_times": [format_time(seconds) for seconds in args.force_time],
        }
        (workdir / "report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(output)
        print(f"pages={len(accepted)} workdir={workdir}")
    except Exception as exc:
        traceback.print_exc()
        print(str(exc), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
